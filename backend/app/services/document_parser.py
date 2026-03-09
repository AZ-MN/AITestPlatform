"""
文档解析服务：支持 PDF、Word、Markdown、TXT、Excel 格式
"""
import os
import re
from typing import Tuple, List, Dict
from pathlib import Path


async def extract_text_from_file(file_path: str, filename: str) -> str:
    """从文件中提取纯文本内容"""
    ext = Path(filename).suffix.lower()

    if ext == ".txt" or ext == ".md":
        return await _read_text_file(file_path)
    elif ext == ".pdf":
        return await _parse_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return await _parse_word(file_path)
    elif ext in (".xlsx", ".xls"):
        return await _parse_excel(file_path)
    else:
        return await _read_text_file(file_path)


async def _read_text_file(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


async def _parse_pdf(file_path: str) -> str:
    try:
        import PyPDF2
        text_parts = []
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        return f"[PDF解析失败: {str(e)}]"


async def _parse_word(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        return f"[Word解析失败: {str(e)}]"


async def _parse_excel(file_path: str) -> str:
    try:
        import pandas as pd
        dfs = pd.read_excel(file_path, sheet_name=None)
        text_parts = []
        for sheet_name, df in dfs.items():
            text_parts.append(f"=== 工作表: {sheet_name} ===")
            text_parts.append(df.to_string(index=False))
        return "\n".join(text_parts)
    except Exception as e:
        return f"[Excel解析失败: {str(e)}]"


def extract_requirement_points_with_rules(content: str) -> List[Dict]:
    """
    基于规则的简单需求点提取（作为AI解析的备用方案）
    """
    points = []
    lines = content.split("\n")
    current_module = "通用"
    point_id = 1

    # 识别模块标题的正则
    module_patterns = [
        r"^#{1,3}\s+(.+)",          # Markdown 标题
        r"^(\d+\.?\s+.{2,30})$",    # 数字编号标题
        r"^[一二三四五六七八九十]+[、.]\s*(.+)$",  # 中文序号
    ]

    # 识别需求描述的关键词
    req_keywords = ["应该", "需要", "必须", "支持", "可以", "允许", "禁止", "不得", "不允许",
                   "shall", "must", "should", "can", "will", "feature", "功能"]

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # 判断是否是模块标题
        is_module = False
        for pattern in module_patterns:
            match = re.match(pattern, line)
            if match:
                text = match.group(1).strip()
                if 2 <= len(text) <= 50:
                    current_module = text
                    is_module = True
                    break

        if not is_module and len(line) > 10:
            # 判断是否包含需求关键词
            has_keyword = any(kw in line for kw in req_keywords)
            if has_keyword or (len(line) > 20 and not line.startswith("-") is False):
                priority = "P1"
                if any(w in line for w in ["核心", "关键", "必须", "must", "critical", "P0"]):
                    priority = "P0"
                elif any(w in line for w in ["可选", "建议", "nice", "P2", "低优"]):
                    priority = "P2"

                points.append({
                    "id": f"REQ-{point_id:03d}",
                    "title": line[:80],
                    "description": line,
                    "priority": priority,
                    "module": current_module,
                    "conditions": [],
                    "rules": []
                })
                point_id += 1

    return points[:100]  # 限制最多100个需求点（规则方式）
